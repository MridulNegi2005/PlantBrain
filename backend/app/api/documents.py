import hashlib
import io
import os
from pathlib import Path

from fastapi import APIRouter, BackgroundTasks, Depends, File, Form, HTTPException, UploadFile
from sqlalchemy.orm import Session

from app.core.config import settings
from app.db import models
from app.db.base import gen_id
from app.db.crud import write_audit
from app.db.seed import PLANT_ID
from app.db.session import get_db
from app.ingestion.pipeline import run_ingestion_bg

router = APIRouter(prefix="/api", tags=["documents"])

UPLOAD_DIR = Path(__file__).resolve().parents[2] / "uploads"

ALLOWED_CONTENT_TYPES = {
    "application/pdf": {".pdf"},
    "image/png": {".png"},
    "image/jpeg": {".jpg", ".jpeg"},
    "text/csv": {".csv"},
    "text/plain": {".txt"},
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": {".xlsx"},
}

INGESTION_STATES = ["uploaded", "extracting", "chunking", "embedding", "graph_building", "completed"]
ACTIVE_INGESTION_STATES = {"queued", "extracting", "chunking", "embedding", "graph_building"}


def _error(status_code: int, code: str, message: str) -> HTTPException:
    return HTTPException(status_code=status_code, detail={"error": {
        "code": code,
        "message": message,
    }})


def _validate_contents(content_type: str, contents: bytes) -> None:
    """Parse enough of an upload to reject empty, corrupt, or spoofed content."""
    if not contents:
        raise _error(400, "empty_file", "Uploaded file is empty.")
    try:
        if content_type == "application/pdf":
            import fitz

            with fitz.open(stream=contents, filetype="pdf") as document:
                if document.page_count < 1:
                    raise ValueError("PDF has no pages")
        elif content_type.startswith("image/"):
            from PIL import Image

            with Image.open(io.BytesIO(contents)) as image:
                image.verify()
        elif content_type.endswith("spreadsheetml.sheet"):
            from openpyxl import load_workbook

            workbook = load_workbook(io.BytesIO(contents), read_only=True, data_only=True)
            workbook.close()
        else:
            decoded = contents.decode("utf-8-sig")
            if "\x00" in decoded:
                raise ValueError("text contains null bytes")
    except Exception as exc:
        if isinstance(exc, HTTPException):
            raise
        raise _error(400, "invalid_file", "File contents do not match the declared type.") from exc


def _doc_out(doc: models.Document) -> dict:
    return {
        "id": doc.id,
        "filename": doc.filename,
        "doc_type": doc.doc_type,
        "status": doc.status,
        "asset_tags": doc.asset_tags or [],
        "created_at": doc.created_at.isoformat(),
    }


@router.post("/documents/upload", status_code=201)
async def upload_document(
    file: UploadFile = File(...),
    plant_id: str = Form(PLANT_ID),
    db: Session = Depends(get_db),
):
    allowed_extensions = ALLOWED_CONTENT_TYPES.get(file.content_type)
    if allowed_extensions is None:
        raise _error(415, "unsupported_media_type", f"Content type {file.content_type} is not allowed.")

    source_suffix = Path(file.filename or "").suffix.lower()
    if source_suffix not in allowed_extensions:
        raise _error(400, "file_type_mismatch", "Filename extension does not match the declared type.")
    if db.get(models.Plant, plant_id) is None:
        raise _error(404, "plant_not_found", f"Plant {plant_id} not found.")

    max_bytes = settings.max_upload_mb * 1024 * 1024
    chunk_size = 1024 * 1024
    contents = bytearray()
    while True:
        chunk = await file.read(chunk_size)
        if not chunk:
            break
        contents.extend(chunk)
        if len(contents) > max_bytes:
            raise _error(413, "file_too_large", f"File exceeds {settings.max_upload_mb}MB limit.")

    _validate_contents(file.content_type, bytes(contents))

    doc_id = gen_id("doc")
    file_hash = hashlib.sha256(contents).hexdigest()

    # Store under a server-generated name (never the client filename) to avoid path
    # traversal; the original name is preserved only as metadata.
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    storage_path = UPLOAD_DIR / f"{doc_id}{source_suffix}"
    storage_path.write_bytes(contents)

    doc = models.Document(
        id=doc_id, plant_id=plant_id, filename=file.filename or f"{doc_id}{source_suffix}",
        doc_type="unclassified", status="uploaded",
        storage_path=str(storage_path), hash_sha256=file_hash, asset_tags=[],
    )
    try:
        db.add(doc)
        write_audit(db, action="document.upload", resource_type="document", resource_id=doc_id)
        db.commit()
    except Exception:
        db.rollback()
        storage_path.unlink(missing_ok=True)
        raise

    return {
        "id": doc.id, "filename": doc.filename, "doc_type": doc.doc_type,
        "status": doc.status, "hash_sha256": doc.hash_sha256,
        "created_at": doc.created_at.isoformat(),
    }


@router.get("/documents")
def list_documents(
    doc_type: str | None = None,
    asset_tag: str | None = None,
    status: str | None = None,
    db: Session = Depends(get_db),
):
    q = db.query(models.Document)
    if doc_type:
        q = q.filter(models.Document.doc_type == doc_type)
    if status:
        q = q.filter(models.Document.status == status)
    docs = q.order_by(models.Document.created_at.desc()).all()
    if asset_tag:
        docs = [d for d in docs if asset_tag in (d.asset_tags or [])]
    return {"items": [_doc_out(d) for d in docs], "total": len(docs)}


@router.get("/documents/{document_id}")
def get_document(document_id: str, db: Session = Depends(get_db)):
    doc = db.get(models.Document, document_id)
    if doc is None:
        raise HTTPException(status_code=404, detail={"error": {
            "code": "not_found", "message": f"Document {document_id} not found.",
        }})
    chunks_count = db.query(models.Chunk).filter(
        models.Chunk.document_id == document_id).count()
    return {**_doc_out(doc), "page_count": doc.page_count, "chunks_count": chunks_count}


@router.get("/documents/{document_id}/chunks")
def get_document_chunks(document_id: str, db: Session = Depends(get_db)):
    """Return persisted evidence chunks, or an explicit empty pending state."""
    doc = db.get(models.Document, document_id)
    if doc is None:
        raise HTTPException(status_code=404, detail={"error": {
            "code": "not_found", "message": f"Document {document_id} not found.",
        }})
    chunks = (
        db.query(models.Chunk)
        .filter(models.Chunk.document_id == document_id)
        .order_by(models.Chunk.page_number)
        .all()
    )
    if chunks and doc.status == "completed":
        items = [{
            "chunk_id": c.id, "page": c.page_number, "text": c.text,
            "bbox": c.bbox, "asset_tags": c.asset_tags or [],
        } for c in chunks]
        return {"items": items, "total": len(items), "stub": False}

    return {"items": [], "total": 0, "stub": doc.status != "completed"}


@router.post("/documents/{document_id}/ingest", status_code=202)
def trigger_ingest(document_id: str, background_tasks: BackgroundTasks,
                   db: Session = Depends(get_db)):
    doc = db.get(models.Document, document_id)
    if doc is None:
        raise HTTPException(status_code=404, detail={"error": {
            "code": "not_found", "message": f"Document {document_id} not found.",
        }})
    active_job = db.query(models.IngestionJob).filter(
        models.IngestionJob.document_id == document_id,
        models.IngestionJob.status.in_(ACTIVE_INGESTION_STATES),
    ).first()
    if active_job is not None:
        raise _error(409, "ingestion_in_progress", f"Document {document_id} is already being ingested.")

    job = models.IngestionJob(id=gen_id("job"), document_id=document_id, status="queued")
    doc.status = "queued"
    db.add(job)
    write_audit(db, action="document.ingest", resource_type="document", resource_id=document_id)
    db.commit()
    # Run the extract -> chunk -> embed pipeline after the response is sent; the
    # frontend polls GET /ingestion-jobs/{id} to watch the status advance.
    background_tasks.add_task(run_ingestion_bg, document_id, job.id, db.get_bind())
    return {"ingestion_job_id": job.id, "status": job.status}


@router.get("/ingestion-jobs/{job_id}")
def get_ingestion_job(job_id: str, db: Session = Depends(get_db)):
    job = db.get(models.IngestionJob, job_id)
    if job is None:
        raise HTTPException(status_code=404, detail={"error": {
            "code": "not_found", "message": f"Ingestion job {job_id} not found.",
        }})
    try:
        idx = INGESTION_STATES.index(job.status)
    except ValueError:
        idx = 0
    return {
        "id": job.id, "document_id": job.document_id, "status": job.status,
        "states": INGESTION_STATES, "current_state_index": idx,
        "error": job.error, "updated_at": job.updated_at.isoformat(),
    }
